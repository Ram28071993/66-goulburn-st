import openpyxl
from datetime import datetime

p = '3. Secure Paper Masterfile/Updated/Secure_Paper_Bin_Tally_PerTenant.xlsx'
wb = openpyxl.load_workbook(p)

SITE = '66 Goulburn Street, Sydney, NSW'

# --- 1. Append new rows to D2 - WD Consolidated (per-tenant block convention, sequential Row ID) ---
ws2 = wb['D2 - WD Consolidated']
max_row_id = 0
for r in range(2, ws2.max_row + 1):
    v = ws2.cell(row=r, column=14).value
    if isinstance(v, (int, float)):
        max_row_id = max(max_row_id, int(v))

new_d2_rows = [
    # date, units, tenant, level, provider, doc
    (datetime(2026, 8, 31), 2, 'Barker Evans', 13, 'Complete Shredding Services', '182691'),
    (datetime(2026, 3, 24), 1, 'Wesfarmers', 10, 'Shred-X', 'JOB-8269940-D8G5'),
    (datetime(2026, 6, 16), 1, 'Wesfarmers', 10, 'Shred-X', 'JOB-8632681-J1R3'),
]

d2_row_ids = {}  # doc -> new Row ID
next_row = ws2.max_row + 1
for date, units, tenant, level, provider, doc in new_d2_rows:
    max_row_id += 1
    row = next_row
    values = [date, SITE, 'Secure paper', 'Mobile bin', 0.24, units, None, tenant, level, provider,
              'No', doc, f'=19.2*0.9*F{row}', max_row_id, 'No Duplication Risk', 'Primary Source',
              'Sufficient', 'Include']
    for c, v in enumerate(values, start=1):
        ws2.cell(row=row, column=c, value=v)
    d2_row_ids[doc] = max_row_id
    next_row += 1

print('New D2 rows added, Row IDs:', d2_row_ids)

# --- 2. Rebuild NABERS Dock 2 block (rows 9..old_max), inserting the 3 new dates in chronological order ---
ws1 = wb['NABERS']
dock1_end = 8  # last Dock 1 row (row 2..8 are Dock 1, confirmed)
old_dock2_start = 9
old_max = ws1.max_row

dock2_rows = []
for r in range(old_dock2_start, old_max + 1):
    vals = [ws1.cell(row=r, column=c).value for c in range(1, 18)]
    dock2_rows.append(vals)

# columns: 1 date,2 premises,3 addr,4 dock,5 wastetype,6 equip,7 size,8 units,9 weight,
# 10 processing,11 rejected,12 blank(version col is header only? actually col12 is 'Version:4' header note, data blank)
# 13 calc weight formula,14 datasource,15 consolidated row ids,16 tenants dock1,17 tenants dock2

new_entries = [
    (datetime(2026, 8, 31), 2, 'Barker Evans', d2_row_ids['182691']),
    (datetime(2026, 3, 24), 1, 'Wesfarmers', d2_row_ids['JOB-8269940-D8G5']),
    (datetime(2026, 6, 16), 1, 'Wesfarmers', d2_row_ids['JOB-8632681-J1R3']),
]

for date, units, tenant, row_id in new_entries:
    vals = [date, 2235, SITE, 'Dock 2', 'Secure paper', 'Mobile bin', 0.24, units, None, None, None, None,
            None, 'Contractor Report, verified no duplication (Method 3 gap-fill)', str(row_id), None,
            f'{tenant}: {units} bin' + ('s' if units != 1 else '')]
    dock2_rows.append(vals)

dock2_rows.sort(key=lambda v: v[0])

# rewrite rows 9.. with sorted data, fixing calc-weight formula to self-reference
new_max = dock1_end + len(dock2_rows)
for i, vals in enumerate(dock2_rows):
    r = dock1_end + 1 + i
    for c in range(1, 18):
        v = vals[c - 1]
        if c == 13:
            v = f'=19.2*0.9*H{r}'
        ws1.cell(row=r, column=c, value=v)

# clear any leftover rows if new_max < old_max (not the case here, but be safe)
if new_max < old_max:
    for r in range(new_max + 1, old_max + 1):
        for c in range(1, 18):
            ws1.cell(row=r, column=c, value=None)

print('NABERS Dock2 block rewritten, rows', dock1_end + 1, 'to', new_max, '(was', old_dock2_start, 'to', old_max, ')')

wb.save(p)
print('saved')
